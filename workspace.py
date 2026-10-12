from __future__ import annotations

import json
import os
import re
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from docx import Document


SECTION_STATUSES = {
    "pending",
    "planned",
    "writing",
    "draft",
    "review",
    "approved",
    "needs_revision",
    "completed",
}

STATUS_LABELS = {
    "pending": "⚪ pending",
    "planned": "🔵 planned",
    "writing": "🟣 writing",
    "draft": "🟡 draft",
    "review": "🔎 review",
    "approved": "✅ approved",
    "needs_revision": "🟠 needs_revision",
    "completed": "✅ completed",
}


class WorkspaceError(RuntimeError):
    pass


@dataclass(slots=True)
class SectionNode:
    id: str
    title: str
    target_pages: float | None = None
    status: str = "pending"
    children: list["SectionNode"] = field(default_factory=list)

    def normalize(self) -> None:
        self.id = str(self.id).strip()
        self.title = str(self.title).strip()
        if self.status not in SECTION_STATUSES:
            self.status = "pending"
        for child in self.children:
            child.normalize()

    @property
    def is_leaf(self) -> bool:
        return not self.children

    def to_dict(self) -> dict[str, Any]:
        return {
            "id": self.id,
            "title": self.title,
            "target_pages": self.target_pages,
            "status": self.status,
            "children": [c.to_dict() for c in self.children],
        }

    @classmethod
    def from_dict(cls, data: dict[str, Any]) -> "SectionNode":
        node = cls(
            id=str(data.get("id") or "").strip(),
            title=str(data.get("title") or "").strip(),
            target_pages=_safe_float(data.get("target_pages")),
            status=str(data.get("status") or "pending"),
            children=[cls.from_dict(x) for x in data.get("children", []) if isinstance(x, dict)],
        )
        node.normalize()
        return node


@dataclass(slots=True)
class WorkPlan:
    title: str
    work_type: str
    target_pages: float | None = None
    sections: list[SectionNode] = field(default_factory=list)
    created_at: str = field(default_factory=lambda: datetime.now().isoformat(timespec="seconds"))
    updated_at: str = field(default_factory=lambda: datetime.now().isoformat(timespec="seconds"))
    schema_version: int = 1

    def to_dict(self) -> dict[str, Any]:
        return {
            "schema_version": self.schema_version,
            "title": self.title,
            "work_type": self.work_type,
            "target_pages": self.target_pages,
            "created_at": self.created_at,
            "updated_at": datetime.now().isoformat(timespec="seconds"),
            "sections": [s.to_dict() for s in self.sections],
        }

    @classmethod
    def from_dict(cls, data: dict[str, Any]) -> "WorkPlan":
        return cls(
            title=str(data.get("title") or "Без темы"),
            work_type=str(data.get("work_type") or "other"),
            target_pages=_safe_float(data.get("target_pages")),
            sections=[SectionNode.from_dict(x) for x in data.get("sections", []) if isinstance(x, dict)],
            created_at=str(data.get("created_at") or datetime.now().isoformat(timespec="seconds")),
            updated_at=str(data.get("updated_at") or datetime.now().isoformat(timespec="seconds")),
            schema_version=int(data.get("schema_version") or 1),
        )

    def walk(self) -> Iterable[SectionNode]:
        def _walk(nodes: list[SectionNode]):
            for node in nodes:
                yield node
                yield from _walk(node.children)
        yield from _walk(self.sections)

    def leaves(self) -> list[SectionNode]:
        return [s for s in self.walk() if s.is_leaf]

    def find(self, section_id: str) -> SectionNode | None:
        sid = normalize_section_id(section_id)
        for node in self.walk():
            if normalize_section_id(node.id) == sid:
                return node
        return None

    def progress(self) -> tuple[int, int, int]:
        leaves = self.leaves()
        done = sum(1 for s in leaves if s.status in {"approved", "completed"})
        total = len(leaves)
        pct = round(done * 100 / total) if total else 0
        return done, total, pct


@dataclass(slots=True)
class ProjectMemory:
    research_object: str | None = None
    research_subject: str | None = None
    goal: str | None = None
    tasks: list[str] = field(default_factory=list)
    terminology: dict[str, str] = field(default_factory=dict)
    organization: str | None = None
    time_period: str | None = None
    accepted_definitions: dict[str, str] = field(default_factory=dict)
    important_numbers: dict[str, Any] = field(default_factory=dict)
    claims: dict[str, str] = field(default_factory=dict)
    updated_at: str = field(default_factory=lambda: datetime.now().isoformat(timespec="seconds"))

    def to_dict(self) -> dict[str, Any]:
        data = asdict(self)
        data["updated_at"] = datetime.now().isoformat(timespec="seconds")
        return data

    @classmethod
    def from_dict(cls, data: dict[str, Any]) -> "ProjectMemory":
        allowed = {f.name for f in cls.__dataclass_fields__.values()}
        return cls(**{k: v for k, v in data.items() if k in allowed})


@dataclass(slots=True)
class SectionVersion:
    version: int
    created_at: str
    reason: str
    source_version: int | None = None
    feedback_file: str | None = None
    word_count: int = 0
    status: str = "draft"
    md_file: str | None = None
    docx_file: str | None = None
    review_file: str | None = None

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


WORK_PLAN_FILE = "work_plan.json"
MEMORY_FILE = "project_memory.json"
SECTIONS_DIR = "sections"


def _safe_float(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _version_keep() -> int:
    """Сколько версий раздела хранить. Переопределяется VERSION_KEEP в config.env."""
    try:
        return max(2, int(os.getenv("VERSION_KEEP", "10")))
    except ValueError:
        return 10


def normalize_section_id(section_id: str) -> str:
    return str(section_id).strip().casefold().replace(" ", "")


def section_slug(section_id: str, title: str = "") -> str:
    sid = normalize_section_id(section_id)
    if sid in {"intro", "introduction", "введение"}:
        return "00_introduction"
    if sid in {"conclusion", "заключение"}:
        return "conclusion"
    if sid in {"references", "literature", "списоклитературы"}:
        return "references"
    cleaned = re.sub(r"[^0-9a-zа-яё]+", "_", sid, flags=re.IGNORECASE).strip("_")
    cleaned = cleaned.replace(".", "_")
    if not cleaned:
        cleaned = "section"
    return cleaned[:70]


def work_plan_path(project: Path) -> Path:
    return project / WORK_PLAN_FILE


def memory_path(project: Path) -> Path:
    return project / MEMORY_FILE


def load_work_plan(project: Path) -> WorkPlan | None:
    path = work_plan_path(project)
    if not path.exists():
        return None
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(data, dict):
            raise ValueError("work_plan.json должен содержать объект")
        return WorkPlan.from_dict(data)
    except Exception as exc:
        raise WorkspaceError(f"Не удалось прочитать work_plan.json: {exc}") from exc


def save_work_plan(project: Path, plan: WorkPlan) -> Path:
    path = work_plan_path(project)
    _atomic_json(path, plan.to_dict())
    ensure_sections_from_plan(project, plan)
    return path


def load_memory(project: Path) -> ProjectMemory:
    path = memory_path(project)
    if not path.exists():
        return ProjectMemory()
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        return ProjectMemory.from_dict(data if isinstance(data, dict) else {})
    except Exception as exc:
        raise WorkspaceError(f"Не удалось прочитать project_memory.json: {exc}") from exc


def save_memory(project: Path, memory: ProjectMemory) -> Path:
    path = memory_path(project)
    _atomic_json(path, memory.to_dict())
    return path


def ensure_sections_from_plan(project: Path, plan: WorkPlan) -> None:
    root = project / SECTIONS_DIR
    root.mkdir(exist_ok=True)
    for node in plan.walk():
        folder = root / section_slug(node.id, node.title)
        folder.mkdir(exist_ok=True)
        metadata = folder / "metadata.json"
        if not metadata.exists():
            _atomic_json(
                metadata,
                {
                    "section_id": node.id,
                    "title": node.title,
                    "current_version": 0,
                    "status": node.status,
                    "versions": [],
                },
            )


def section_dir(project: Path, section: SectionNode | str, title: str = "") -> Path:
    if isinstance(section, SectionNode):
        sid, title = section.id, section.title
    else:
        sid = section
    path = project / SECTIONS_DIR / section_slug(sid, title)
    path.mkdir(parents=True, exist_ok=True)
    return path


def load_section_metadata(folder: Path, *, section_id: str = "", title: str = "") -> dict[str, Any]:
    path = folder / "metadata.json"
    if path.exists():
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
            if isinstance(data, dict):
                data.setdefault("versions", [])
                data.setdefault("current_version", 0)
                data.setdefault("status", "pending")
                return data
        except Exception as exc:
            raise WorkspaceError(f"Не удалось прочитать {path}: {exc}") from exc
    data = {
        "section_id": section_id,
        "title": title,
        "current_version": 0,
        "status": "pending",
        "versions": [],
    }
    _atomic_json(path, data)
    return data


def save_section_metadata(folder: Path, data: dict[str, Any]) -> None:
    _atomic_json(folder / "metadata.json", data)


def latest_section_version(project: Path, node: SectionNode) -> tuple[SectionVersion | None, str | None]:
    folder = section_dir(project, node)
    meta = load_section_metadata(folder, section_id=node.id, title=node.title)
    versions = [x for x in meta.get("versions", []) if isinstance(x, dict)]
    if not versions:
        return None, None
    current = int(meta.get("current_version") or max(int(v.get("version") or 0) for v in versions))
    entry = next((v for v in versions if int(v.get("version") or 0) == current), versions[-1])
    version = SectionVersion(**{k: entry.get(k) for k in SectionVersion.__dataclass_fields__.keys()})
    text = None
    if version.md_file:
        p = folder / version.md_file
        if p.exists():
            text = p.read_text(encoding="utf-8")
    return version, text


def save_section_version(
    project: Path,
    node: SectionNode,
    text: str,
    *,
    reason: str = "generation",
    source_version: int | None = None,
    feedback_file: str | None = None,
    status: str = "draft",
    docx_builder=None,
    review: dict[str, Any] | None = None,
) -> SectionVersion:
    folder = section_dir(project, node)
    meta = load_section_metadata(folder, section_id=node.id, title=node.title)
    current = int(meta.get("current_version") or 0)
    number = current + 1
    md_name = f"v{number}.md"
    docx_name = f"v{number}.docx"
    review_name = f"v{number}_review.json" if review else None
    (folder / md_name).write_text(text.strip() + "\n", encoding="utf-8")
    if docx_builder:
        docx_builder(folder / docx_name, text)
    else:
        _minimal_docx(folder / docx_name, node.title, text)
    if review:
        _atomic_json(folder / str(review_name), review)
    version = SectionVersion(
        version=number,
        created_at=datetime.now().isoformat(timespec="seconds"),
        reason=reason,
        source_version=source_version,
        feedback_file=feedback_file,
        word_count=len(re.findall(r"\S+", text)),
        status=status if status in SECTION_STATUSES else "draft",
        md_file=md_name,
        docx_file=docx_name,
        review_file=review_name,
    )
    versions = [x for x in meta.get("versions", []) if isinstance(x, dict)]
    versions.append(version.to_dict())
    # Ротация версий: храним не больше VERSION_KEEP последних, старые файлы удаляем,
    # иначе sections/ и metadata.json разрастаются бесконечно.
    keep = _version_keep()
    if len(versions) > keep:
        excess = versions[:-keep]
        for old in excess:
            for key in ("md_file", "docx_file", "review_file"):
                old_name = old.get(key)
                if old_name:
                    try:
                        (folder / str(old_name)).unlink(missing_ok=True)
                    except OSError:
                        pass
        versions = versions[-keep:]
    meta["versions"] = versions
    meta["current_version"] = number
    meta["status"] = version.status
    save_section_metadata(folder, meta)
    node.status = version.status
    _update_plan_node_status(project, node.id, version.status)
    return version


def set_section_status(project: Path, node: SectionNode, status: str) -> None:
    if status not in SECTION_STATUSES:
        raise WorkspaceError(f"Неизвестный статус раздела: {status}")
    folder = section_dir(project, node)
    meta = load_section_metadata(folder, section_id=node.id, title=node.title)
    meta["status"] = status
    current = int(meta.get("current_version") or 0)
    for item in meta.get("versions", []):
        if int(item.get("version") or 0) == current:
            item["status"] = status
    save_section_metadata(folder, meta)
    node.status = status
    _update_plan_node_status(project, node.id, status)


def list_section_versions(project: Path, node: SectionNode) -> list[SectionVersion]:
    folder = section_dir(project, node)
    meta = load_section_metadata(folder, section_id=node.id, title=node.title)
    result: list[SectionVersion] = []
    for item in meta.get("versions", []):
        if not isinstance(item, dict):
            continue
        values = {k: item.get(k) for k in SectionVersion.__dataclass_fields__.keys()}
        result.append(SectionVersion(**values))
    return result


def save_section_summary(project: Path, node: SectionNode, summary: str) -> Path:
    path = section_dir(project, node) / "summary.md"
    path.write_text(summary.strip() + "\n", encoding="utf-8")
    return path


def load_section_summary(project: Path, node: SectionNode) -> str:
    path = section_dir(project, node) / "summary.md"
    return path.read_text(encoding="utf-8").strip() if path.exists() else ""


def previous_summaries(project: Path, plan: WorkPlan, current_id: str, max_chars: int = 18_000) -> str:
    pieces: list[str] = []
    total = 0
    for node in plan.leaves():
        if normalize_section_id(node.id) == normalize_section_id(current_id):
            break
        summary = load_section_summary(project, node)
        if not summary:
            continue
        block = f"\n### {node.id} {node.title}\n{summary}\n"
        if total + len(block) > max_chars:
            break
        pieces.append(block)
        total += len(block)
    return "".join(pieces).strip()


def plan_text(plan: WorkPlan) -> str:
    done, total, pct = plan.progress()
    lines = [
        f"📚 {plan.title}",
        f"Готовность: {pct}% ({done}/{total})",
        "",
    ]

    def add(nodes: list[SectionNode], level: int = 0) -> None:
        for node in nodes:
            prefix = "   " * level
            label = STATUS_LABELS.get(node.status, node.status)
            pages = f" · ~{node.target_pages:g} стр." if node.target_pages else ""
            lines.append(f"{prefix}{node.id} {node.title}{pages} — {label}")
            add(node.children, level + 1)

    add(plan.sections)
    return "\n".join(lines)


def default_work_plan(title: str, work_type: str, target_pages: float | None = None) -> WorkPlan:
    pages = target_pages or (35 if work_type in {"coursework", "thesis"} else None)
    if work_type in {"coursework", "thesis"}:
        sections = [
            SectionNode("intro", "Введение", 2),
            SectionNode("1", "Теоретические основы исследования", 12, children=[
                SectionNode("1.1", "Понятия и теоретические подходы", 4),
                SectionNode("1.2", "Ключевые модели и методы", 4),
                SectionNode("1.3", "Современные подходы и промежуточные выводы", 4),
            ]),
            SectionNode("2", "Аналитическая часть", 14, children=[
                SectionNode("2.1", "Характеристика объекта исследования", 4),
                SectionNode("2.2", "Анализ текущего состояния", 5),
                SectionNode("2.3", "Проблемы и направления совершенствования", 5),
            ]),
            SectionNode("conclusion", "Заключение", 2),
            SectionNode("references", "Список литературы", 2),
        ]
    else:
        sections = [
            SectionNode("1", "Основная часть", None),
            SectionNode("conclusion", "Вывод", None),
        ]
    return WorkPlan(title=title or "Без темы", work_type=work_type, target_pages=pages, sections=sections)


def _update_plan_node_status(project: Path, section_id: str, status: str) -> None:
    plan = load_work_plan(project)
    if plan is None:
        return
    node = plan.find(section_id)
    if node is None:
        return
    node.status = status
    save_work_plan(project, plan)


def _atomic_json(path: Path, data: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(path)


def _minimal_docx(path: Path, title: str, text: str) -> None:
    doc = Document()
    doc.add_heading(title, level=1)
    for raw in text.splitlines():
        line = raw.strip()
        if line:
            doc.add_paragraph(line)
    doc.save(path)
