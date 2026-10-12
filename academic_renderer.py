from __future__ import annotations

import re
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Sequence

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


@dataclass(slots=True)
class FormattingProfile:
    font_name: str = "Times New Roman"
    font_size_pt: float = 14.0
    heading1_size_pt: float = 17.0
    heading2_size_pt: float = 16.0
    heading3_size_pt: float = 14.0
    line_spacing: float = 1.5
    first_line_indent_cm: float = 1.25
    margin_left_cm: float = 3.0
    margin_right_cm: float = 1.5
    margin_top_cm: float = 2.0
    margin_bottom_cm: float = 2.0
    alignment: str = "justify"

    def to_dict(self) -> dict:
        return asdict(self)

    @classmethod
    def from_dict(cls, data: dict) -> "FormattingProfile":
        base = cls()
        for key in cls.__dataclass_fields__:
            if key in data and data[key] not in (None, ""):
                setattr(base, key, data[key])
        base._normalize()
        return base

    def _normalize(self) -> None:
        self.font_name = str(self.font_name or "Times New Roman")[:80]
        for attr, low, high in [
            ("font_size_pt", 8, 24),
            ("heading1_size_pt", 10, 28),
            ("heading2_size_pt", 10, 26),
            ("heading3_size_pt", 10, 24),
            ("line_spacing", 1, 3),
            ("first_line_indent_cm", 0, 4),
            ("margin_left_cm", 0.5, 6),
            ("margin_right_cm", 0.5, 6),
            ("margin_top_cm", 0.5, 6),
            ("margin_bottom_cm", 0.5, 6),
        ]:
            try:
                value = float(getattr(self, attr))
            except Exception:
                value = float(getattr(FormattingProfile(), attr))
            setattr(self, attr, min(max(value, low), high))
        if self.alignment not in {"justify", "left", "right", "center"}:
            self.alignment = "justify"


class AcademicDocxRenderer:
    """Детерминированный DOCX renderer: методичка задаёт профиль, иначе AcademicPreset."""

    def __init__(self, profile: FormattingProfile | None = None) -> None:
        self.profile = profile or FormattingProfile()
        self.profile._normalize()

    def render(
        self,
        path: Path,
        text: str,
        *,
        title: str | None = None,
        graph_paths: Sequence[Path] = (),
        source_number_map: dict[str, int] | None = None,
    ) -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        doc = Document()
        self._setup_doc(doc)
        if title:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(title)
            run.bold = True
            run.font.name = self.profile.font_name
            run.font.size = Pt(self.profile.heading1_size_pt)
            self._black(run)

        lines = text.splitlines()
        for raw in lines:
            line = raw.rstrip()
            stripped = line.strip()
            if not stripped:
                doc.add_paragraph("")
                continue

            if source_number_map:
                stripped = self._replace_source_ids(stripped, source_number_map)

            if stripped.startswith("### "):
                p = self._heading(doc, stripped[4:].strip(), 3)
            elif stripped.startswith("## "):
                p = self._heading(doc, stripped[3:].strip(), 2)
            elif stripped.startswith("# "):
                p = self._heading(doc, stripped[2:].strip(), 1)
            elif re.match(r"^[-•]\s+", stripped):
                p = doc.add_paragraph(re.sub(r"^[-•]\s+", "", stripped), style="List Bullet")
                self._paragraph_style(p, first_indent=False)
            elif re.match(r"^\d+[.)]\s+", stripped):
                p = doc.add_paragraph(re.sub(r"^\d+[.)]\s+", "", stripped), style="List Number")
                self._paragraph_style(p, first_indent=False)
            else:
                p = doc.add_paragraph(stripped)
                self._paragraph_style(p, first_indent=True)

            for run in p.runs:
                run.font.name = self.profile.font_name
                if not run.font.size:
                    run.font.size = Pt(self.profile.font_size_pt)
                self._black(run)

        for graph in graph_paths:
            if not graph.exists():
                continue
            try:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(graph), width=Cm(15.5))
                caption = doc.add_paragraph(graph.stem.replace("_", " "))
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self._paragraph_style(caption, first_indent=False)
            except Exception:
                continue

        doc.save(path)
        return path

    def render_combined(
        self,
        path: Path,
        sections: Sequence[tuple[str, str]],
        *,
        title: str | None = None,
        source_number_map: dict[str, int] | None = None,
    ) -> Path:
        blocks: list[str] = []
        for heading, text in sections:
            if heading:
                blocks.append(f"# {heading}")
            blocks.append(text.strip())
        return self.render(path, "\n\n".join(blocks), title=title, source_number_map=source_number_map)

    def _setup_doc(self, doc: Document) -> None:
        section = doc.sections[0]
        section.left_margin = Cm(self.profile.margin_left_cm)
        section.right_margin = Cm(self.profile.margin_right_cm)
        section.top_margin = Cm(self.profile.margin_top_cm)
        section.bottom_margin = Cm(self.profile.margin_bottom_cm)
        normal = doc.styles["Normal"]
        normal.font.name = self.profile.font_name
        normal.font.size = Pt(self.profile.font_size_pt)

        for name, size in [("Heading 1", self.profile.heading1_size_pt), ("Heading 2", self.profile.heading2_size_pt), ("Heading 3", self.profile.heading3_size_pt)]:
            style = doc.styles[name]
            style.font.name = self.profile.font_name
            style.font.size = Pt(size)
            style.font.bold = True
            # Цвет в style XML — чёрный.
            rpr = style.element.get_or_add_rPr()
            color = rpr.find(qn("w:color"))
            if color is None:
                color = OxmlElement("w:color")
                rpr.append(color)
            color.set(qn("w:val"), "000000")

    def _heading(self, doc: Document, text: str, level: int):
        p = doc.add_heading(text, level=level)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = self.profile.line_spacing
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        size = {1: self.profile.heading1_size_pt, 2: self.profile.heading2_size_pt, 3: self.profile.heading3_size_pt}[level]
        for run in p.runs:
            run.font.name = self.profile.font_name
            run.font.size = Pt(size)
            run.bold = True
            self._black(run)
        return p

    def _paragraph_style(self, p, *, first_indent: bool) -> None:
        mapping = {
            "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
        }
        p.alignment = mapping.get(self.profile.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
        fmt = p.paragraph_format
        fmt.first_line_indent = Cm(self.profile.first_line_indent_cm if first_indent else 0)
        fmt.line_spacing = self.profile.line_spacing
        fmt.space_after = Pt(0)
        fmt.space_before = Pt(0)

    @staticmethod
    def _black(run) -> None:
        color = run._r.get_or_add_rPr().find(qn("w:color"))
        if color is None:
            color = OxmlElement("w:color")
            run._r.get_or_add_rPr().append(color)
        color.set(qn("w:val"), "000000")

    @staticmethod
    def _replace_source_ids(text: str, source_number_map: dict[str, int]) -> str:
        def repl(match: re.Match) -> str:
            sid = match.group(1).upper()
            number = source_number_map.get(sid)
            return f"[{number}]" if number else match.group(0)
        return re.sub(r"\[(SOURCE_\d{3,})\]", repl, text, flags=re.IGNORECASE)
