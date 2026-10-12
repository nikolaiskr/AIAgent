from __future__ import annotations

import argparse
import json
import logging
import math
import re
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Sequence

from docx import Document
from docx.shared import Cm, Pt

from projects import (
    ProjectError,
    create_project,
    get_project,
    list_materials,
    list_projects,
    load_project_meta,
    materials_text,
    project_summary,
    read_materials,
    readiness_text,
    safe_name,
)
from providers import HybridRouter, ProviderError
from readers import SourceDocument
from settings import LOG_PATH, Settings, load_settings
from pipeline import create_work_plan, generate_chapter, generate_section
from assembler import assemble_work
from workspace import load_work_plan, plan_text

try:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import numpy as np
    import sympy as sp
except Exception:
    plt = None
    np = None
    sp = None


LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(name)s | %(message)s",
    handlers=[logging.FileHandler(LOG_PATH, encoding="utf-8"), logging.StreamHandler(sys.stdout)],
)
logger = logging.getLogger("aiagent")


class AgentError(RuntimeError):
    pass


@dataclass(slots=True)
class Chunk:
    source_name: str
    text: str
    score: float = 0.0


@dataclass(slots=True)
class SolveResult:
    answer: str
    docx_path: Path
    txt_path: Path
    graph_paths: list[Path] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


@dataclass(slots=True)
class FormattingProfile:
    font_name: str = "Times New Roman"
    font_size_pt: float = 14.0
    line_spacing: float = 1.5
    first_line_indent_cm: float = 1.25
    margin_left_cm: float = 3.0
    margin_right_cm: float = 1.5
    margin_top_cm: float = 2.0
    margin_bottom_cm: float = 2.0
    alignment: str = "justify"


def _tokenize(text: str) -> set[str]:
    tokens = re.findall(r"[a-zа-яё0-9]{2,}", text.casefold())
    stop = {
        "это", "как", "для", "что", "или", "при", "над", "под", "все", "всех", "где",
        "нужно", "вариант", "задание", "задача", "the", "and", "with", "from",
    }
    return {t for t in tokens if t not in stop}


def chunk_documents(documents: Sequence[SourceDocument], settings: Settings) -> list[Chunk]:
    chunks: list[Chunk] = []
    step = max(1, settings.chunk_chars - settings.chunk_overlap)
    for doc in documents:
        text = doc.text.strip()
        if not text:
            continue
        if len(text) <= settings.chunk_chars:
            chunks.append(Chunk(source_name=doc.path.name, text=text))
            continue
        for start in range(0, len(text), step):
            part = text[start:start + settings.chunk_chars].strip()
            if part:
                chunks.append(Chunk(source_name=doc.path.name, text=part))
            if start + settings.chunk_chars >= len(text):
                break
    return chunks


def select_context(documents: Sequence[SourceDocument], query: str, settings: Settings, *, formatting_only: bool = False) -> str:
    chunks = chunk_documents(documents, settings)
    if not chunks:
        return "[Исходные документы не содержат извлекаемого текста.]"

    query_tokens = _tokenize(query)
    query_lower = query.casefold()
    numbers = set(re.findall(r"\d+(?:[.,]\d+)?", query_lower))
    method_hints = ("метод", "оформ", "требован", "гост", "указан", "рекомендац", "регламент", "пример")

    for chunk in chunks:
        lower = chunk.text.casefold()
        tokens = _tokenize(chunk.text)
        overlap = len(query_tokens & tokens)
        numeric_hits = sum(1 for n in numbers if n in lower)
        method_bonus = 3 if any(h in chunk.source_name.casefold() for h in method_hints) else 0
        keyword_hits = 0
        if formatting_only:
            method_bonus *= 4
            keyword_hits = sum(
                1 for kw in ("шрифт", "поле", "интервал", "абзац", "оформ", "литератур", "титуль")
                if kw in lower
            )
        exact_bonus = 8 if query_lower and query_lower[:120] in lower else 0
        chunk.score = overlap * 1.5 + numeric_hits * 4 + method_bonus + keyword_hits * 2 + exact_bonus

    ranked = sorted(chunks, key=lambda c: c.score, reverse=True)
    chosen = ranked[:settings.top_chunks]
    blocks: list[str] = []
    total = 0
    for chunk in chosen:
        block = f"\n===== ИСТОЧНИК: {chunk.source_name} =====\n{chunk.text}\n"
        if total + len(block) > settings.max_context_chars:
            remaining = settings.max_context_chars - total
            if remaining > 800:
                blocks.append(block[:remaining])
            break
        blocks.append(block)
        total += len(block)
    return "".join(blocks).strip()


def read_soul(settings: Settings) -> str:
    candidates = [settings.projects_root / "soul.md", Path(__file__).resolve().parent / "soul.md"]
    for path in candidates:
        if path.exists() and path.is_file():
            for encoding in ("utf-8-sig", "utf-8", "cp1251"):
                try:
                    return path.read_text(encoding=encoding).strip()
                except UnicodeDecodeError:
                    continue
                except OSError:
                    break
    return (
        "Ты — аккуратный учебный ассистент. Не выдумывай данные и источники, проверяй вычисления, "
        "явно сообщай о недостатке исходных данных и соблюдай методические требования."
    )


def build_solution_messages(user_request: str, context: str, soul: str) -> list[dict[str, str]]:
    system = f"""
{soul}

Обязательные правила:
1. Используй только материалы, действительно относящиеся к запросу.
2. Если данных не хватает — перечисли, чего именно не хватает; не выдумывай условие.
3. Для вычислений показывай ход решения и проверку результата.
4. Не придумывай библиографические источники.
5. Текст должен быть пригоден для переноса в DOCX.
6. Если нужен график функции от x, добавь после текста блок:
<<<GRAPH>>>
{{"title":"Название","expression":"x**2 - 4*x + 3","x_min":-5,"x_max":8,"x_label":"x","y_label":"y"}}
<<<END_GRAPH>>>
""".strip()
    user = f"""
ЗАПРОС:
{user_request}

МАТЕРИАЛЫ:
{context}

Подготовь результат с учётом материалов и требований.
""".strip()
    return [{"role": "system", "content": system}, {"role": "user", "content": user}]


def extract_graph_specs(answer: str) -> tuple[str, list[dict]]:
    pattern = re.compile(r"<<<GRAPH>>>\s*(\{.*?\})\s*<<<END_GRAPH>>>", re.DOTALL | re.IGNORECASE)
    specs: list[dict] = []
    for match in pattern.finditer(answer):
        try:
            specs.append(json.loads(match.group(1)))
        except json.JSONDecodeError:
            logger.warning("Некорректный GRAPH JSON")
    return pattern.sub("", answer).strip(), specs


def _bounded_float(value: object, default: float, low: float, high: float) -> float:
    try:
        num = float(value)
        if math.isfinite(num) and low <= num <= high:
            return num
    except (TypeError, ValueError):
        pass
    return default


def _validate_expression(expr: str) -> bool:
    if len(expr) > 300 or re.search(r"[_\[\]{}'\";:]", expr):
        return False
    if not re.fullmatch(r"[0-9a-zA-Zа-яА-ЯёЁ+\-*/().,^\s]+", expr):
        return False
    words = set(re.findall(r"[A-Za-z]+", expr))
    return words <= {"x", "sin", "cos", "tan", "exp", "log", "sqrt", "abs", "pi", "e"}


def render_graphs(specs: Sequence[dict], output_dir: Path, base_name: str) -> tuple[list[Path], list[str]]:
    if not specs:
        return [], []
    if plt is None or np is None or sp is None:
        return [], ["График запрошен, но matplotlib/numpy/sympy не импортировались."]

    paths: list[Path] = []
    warnings: list[str] = []
    x = sp.symbols("x", real=True)
    locals_map = {"x": x, "sin": sp.sin, "cos": sp.cos, "tan": sp.tan, "exp": sp.exp,
                  "log": sp.log, "sqrt": sp.sqrt, "abs": sp.Abs, "pi": sp.pi, "e": sp.E}

    for idx, spec in enumerate(specs, start=1):
        try:
            expr_text = str(spec.get("expression", "")).replace("^", "**").strip()
            if not _validate_expression(expr_text):
                raise ValueError("выражение не прошло безопасную проверку")
            expr = sp.sympify(expr_text, locals=locals_map)
            if expr.free_symbols - {x}:
                raise ValueError("разрешена только переменная x")
            x_min = _bounded_float(spec.get("x_min"), -10, -1_000_000, 1_000_000)
            x_max = _bounded_float(spec.get("x_max"), 10, -1_000_000, 1_000_000)
            if x_max <= x_min:
                raise ValueError("x_max должен быть больше x_min")
            fn = sp.lambdify(x, expr, modules=["numpy"])
            xs = np.linspace(x_min, x_max, 1200)
            with np.errstate(all="ignore"):
                ys = np.asarray(fn(xs), dtype=float)
            if ys.ndim == 0:
                ys = np.full_like(xs, float(ys))
            ys[~np.isfinite(ys)] = np.nan
            fig = plt.figure(figsize=(8, 5))
            ax = fig.add_subplot(111)
            ax.plot(xs, ys)
            ax.axhline(0, linewidth=0.8)
            ax.axvline(0, linewidth=0.8)
            ax.grid(True, alpha=0.25)
            ax.set_title(str(spec.get("title") or f"График {idx}"))
            ax.set_xlabel(str(spec.get("x_label") or "x"))
            ax.set_ylabel(str(spec.get("y_label") or "y"))
            fig.tight_layout()
            path = output_dir / f"{base_name}_график_{idx}.png"
            fig.savefig(path, dpi=180)
            plt.close(fig)
            paths.append(path)
        except Exception as exc:
            warnings.append(f"Не удалось построить график {idx}: {exc}")
    return paths, warnings


def _format_profile_from_ai(settings: Settings, documents: Sequence[SourceDocument]) -> FormattingProfile:
    router = HybridRouter(settings)
    provider = router.quality()
    if provider is None:
        return FormattingProfile()
    context = select_context(
        documents,
        "правила оформления шрифт размер поля интервал абзац",
        settings,
        formatting_only=True,
    )
    prompt = f"""
Извлеки только параметры оформления основной части из материалов. Верни только JSON:
{{"font_name":"Times New Roman","font_size_pt":14,"line_spacing":1.5,"first_line_indent_cm":1.25,
"margin_left_cm":3.0,"margin_right_cm":1.5,"margin_top_cm":2.0,"margin_bottom_cm":2.0,"alignment":"justify"}}
Если параметр не найден — используй указанное значение по умолчанию.

МАТЕРИАЛЫ:
{context[:40_000]}
""".strip()
    try:
        raw = provider.complete([{"role": "user", "content": prompt}], max_tokens=900, temperature=0)
        start, end = raw.find("{"), raw.rfind("}")
        data = json.loads(raw[start:end+1])
        profile = FormattingProfile()
        profile.font_name = str(data.get("font_name") or profile.font_name)[:80]
        profile.font_size_pt = _bounded_float(data.get("font_size_pt"), 14, 8, 24)
        profile.line_spacing = _bounded_float(data.get("line_spacing"), 1.5, 1, 3)
        profile.first_line_indent_cm = _bounded_float(data.get("first_line_indent_cm"), 1.25, 0, 4)
        profile.margin_left_cm = _bounded_float(data.get("margin_left_cm"), 3, .5, 6)
        profile.margin_right_cm = _bounded_float(data.get("margin_right_cm"), 1.5, .5, 6)
        profile.margin_top_cm = _bounded_float(data.get("margin_top_cm"), 2, .5, 6)
        profile.margin_bottom_cm = _bounded_float(data.get("margin_bottom_cm"), 2, .5, 6)
        alignment = str(data.get("alignment") or "justify").casefold()
        profile.alignment = alignment if alignment in {"justify", "left", "center", "right"} else "justify"
        return profile
    except Exception as exc:
        logger.warning("Не удалось определить профиль оформления через AI: %s", exc)
        return FormattingProfile()


def save_solution_docx(path: Path, answer: str, profile: FormattingProfile, graph_paths: Sequence[Path]) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    align_map = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }
    doc = Document()
    section = doc.sections[0]
    section.left_margin = Cm(profile.margin_left_cm)
    section.right_margin = Cm(profile.margin_right_cm)
    section.top_margin = Cm(profile.margin_top_cm)
    section.bottom_margin = Cm(profile.margin_bottom_cm)
    normal = doc.styles["Normal"]
    normal.font.name = profile.font_name
    normal.font.size = Pt(profile.font_size_pt)

    for line in answer.splitlines():
        stripped = line.strip()
        if not stripped:
            doc.add_paragraph("")
            continue
        if stripped.startswith("### "):
            p = doc.add_heading(stripped[4:].strip(), level=3)
        elif stripped.startswith("## "):
            p = doc.add_heading(stripped[3:].strip(), level=2)
        elif stripped.startswith("# "):
            p = doc.add_heading(stripped[2:].strip(), level=1)
        elif re.match(r"^[-•]\s+", stripped):
            p = doc.add_paragraph(re.sub(r"^[-•]\s+", "", stripped), style="List Bullet")
        elif re.match(r"^\d+[.)]\s+", stripped):
            p = doc.add_paragraph(re.sub(r"^\d+[.)]\s+", "", stripped), style="List Number")
        else:
            p = doc.add_paragraph(stripped)
        p.alignment = align_map.get(profile.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
        p.paragraph_format.first_line_indent = Cm(profile.first_line_indent_cm)
        p.paragraph_format.line_spacing = profile.line_spacing
        p.paragraph_format.space_after = Pt(0)
        for run in p.runs:
            run.font.name = profile.font_name
            run.font.size = Pt(profile.font_size_pt)

    for graph in graph_paths:
        try:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(graph), width=Cm(15.5))
            caption = doc.add_paragraph(graph.stem.replace("_", " "))
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception as exc:
            logger.warning("Не удалось вставить график %s: %s", graph, exc)
    doc.save(str(path))


def solve_task(settings: Settings, project_name: str, user_request: str) -> SolveResult:
    if not user_request.strip():
        raise AgentError("Пустой запрос.")
    if not settings.any_ai_ready:
        raise AgentError(
            "AI-движок пока не подключён. Локальная часть работает; для выполнения заполни ключ и модель OpenAI или DeepSeek."
        )

    project = get_project(settings, project_name)
    documents, warnings = read_materials(settings, project_name)
    context = select_context(documents, user_request, settings)
    soul = read_soul(settings)
    router = HybridRouter(settings)

    try:
        writer = router.writer()
        answer_raw = writer.complete(
            build_solution_messages(user_request, context, soul),
            max_tokens=settings.max_output_tokens,
            temperature=0.15,
        )
    except ProviderError as exc:
        raise AgentError(str(exc)) from exc

    answer, graph_specs = extract_graph_specs(answer_raw)

    # Если доступны оба провайдера, OpenAI выступает как финальный quality-review.
    quality = router.quality()
    if settings.deepseek_ready and quality is not None:
        review_prompt = f"""
Проведи финальную редактуру результата учебной работы. Исправь только реальные проблемы: несоответствие запросу,
противоречия, плохую структуру, неясные формулировки. Не добавляй выдуманных фактов или источников.
Верни полностью готовый исправленный текст без комментариев редактора.

ЗАПРОС:
{user_request}

РЕЗУЛЬТАТ:
{answer}
""".strip()
        try:
            answer = quality.complete(
                [{"role": "system", "content": soul}, {"role": "user", "content": review_prompt}],
                max_tokens=settings.max_output_tokens,
                temperature=0.1,
            )
        except Exception as exc:
            warnings.append(f"Финальный OpenAI-review не выполнен: {exc}")

    output_dir = project / "output"
    output_dir.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    base = safe_name(f"Результат_{user_request}", max_len=62) + f"_{stamp}"
    graph_paths, graph_warnings = render_graphs(graph_specs, output_dir, base)
    warnings.extend(graph_warnings)
    profile = _format_profile_from_ai(settings, documents)
    txt_path = output_dir / f"{base}.txt"
    docx_path = output_dir / f"{base}.docx"
    try:
        txt_path.write_text(answer, encoding="utf-8")
        save_solution_docx(docx_path, answer, profile, graph_paths)
    except Exception as exc:
        raise AgentError(f"Не удалось сохранить результат: {exc}") from exc
    return SolveResult(answer=answer, docx_path=docx_path, txt_path=txt_path, graph_paths=graph_paths, warnings=warnings)


def list_docx_files(settings: Settings, project_name: str) -> list[Path]:
    project = get_project(settings, project_name)
    files = [p for p in project.rglob("*.docx") if p.is_file() and not p.name.startswith("~$")]
    return sorted(files, key=lambda p: p.stat().st_mtime, reverse=True)


def inspect_docx_format(path: Path) -> dict:
    try:
        doc = Document(str(path))
    except Exception as exc:
        raise AgentError(f"Не удалось открыть DOCX: {exc}") from exc
    section = doc.sections[0]
    normal = doc.styles["Normal"]
    non_empty = [p for p in doc.paragraphs if p.text.strip()]
    run_fonts: dict[str, int] = {}
    run_sizes: dict[str, int] = {}
    for p in non_empty[:1000]:
        for run in p.runs:
            if run.font.name:
                run_fonts[run.font.name] = run_fonts.get(run.font.name, 0) + 1
            if run.font.size:
                key = str(round(run.font.size.pt, 1))
                run_sizes[key] = run_sizes.get(key, 0) + 1

    def cm(value) -> float | None:
        return round(value.cm, 2) if value else None

    return {
        "file": path.name,
        "paragraphs": len(non_empty),
        "tables": len(doc.tables),
        "images": len(doc.inline_shapes),
        "normal_font": normal.font.name,
        "normal_size": round(normal.font.size.pt, 1) if normal.font.size else None,
        "margins_cm": {
            "left": cm(section.left_margin), "right": cm(section.right_margin),
            "top": cm(section.top_margin), "bottom": cm(section.bottom_margin),
        },
        "run_fonts": run_fonts,
        "run_sizes": run_sizes,
    }


def local_formatting_report(path: Path) -> str:
    data = inspect_docx_format(path)
    default = FormattingProfile()
    margins = data["margins_cm"]
    lines = [
        f"Локальная техническая проверка: {data['file']}",
        "AI и методичка в этой проверке не используются.",
        "",
        f"Абзацев: {data['paragraphs']}",
        f"Таблиц: {data['tables']}",
        f"Изображений: {data['images']}",
        f"Normal font: {data['normal_font'] or 'не задан'}",
        f"Normal size: {data['normal_size'] or 'не задан'}",
        f"Поля, см: {margins}",
        f"Шрифты в runs: {json.dumps(data['run_fonts'], ensure_ascii=False)}",
        f"Размеры в runs: {json.dumps(data['run_sizes'], ensure_ascii=False)}",
        "",
        "Ориентир по умолчанию (НЕ требования конкретного вуза):",
        f"Times New Roman {default.font_size_pt:g}, интервал {default.line_spacing:g}, поля {default.margin_left_cm:g}/{default.margin_right_cm:g}/{default.margin_top_cm:g}/{default.margin_bottom_cm:g} см.",
    ]
    return "\n".join(lines)


def check_formatting(settings: Settings, project_name: str, document_name: str) -> tuple[str, Path]:
    project = get_project(settings, project_name)
    target = (project / document_name).resolve()
    try:
        target.relative_to(project.resolve())
    except ValueError as exc:
        raise AgentError("Недопустимый путь к документу.") from exc
    if not target.exists() or target.suffix.lower() != ".docx":
        raise AgentError(f"DOCX не найден: {document_name}")

    local_report = local_formatting_report(target)
    report = local_report

    if settings.openai_ready:
        documents, warnings = read_materials(settings, project_name)
        method_context = select_context(
            documents,
            "чек-лист оформления методические требования поля шрифт интервал нумерация заголовки литература",
            settings,
            formatting_only=True,
        )
        router = HybridRouter(settings)
        quality = router.quality()
        if quality is not None:
            prompt = f"""
Проверь оформление DOCX по методическим материалам. Не придумывай требований, которых нет.
Верни практичный чек-лист: ✅ соответствует / ⚠️ проверить вручную / ❌ не соответствует / что исправить.

ЛОКАЛЬНАЯ ТЕХНИЧЕСКАЯ ДИАГНОСТИКА:
{local_report}

МЕТОДИЧЕСКИЕ МАТЕРИАЛЫ:
{method_context[:45_000]}

Предупреждения: {warnings[:10]}
""".strip()
            try:
                report = quality.complete([{"role": "user", "content": prompt}], max_tokens=4500, temperature=0.1)
            except Exception as exc:
                report = local_report + f"\n\n⚠️ AI-проверка не выполнена: {exc}"

    output_dir = project / "output"
    output_dir.mkdir(exist_ok=True)
    report_path = output_dir / f"Проверка_оформления_{safe_name(target.stem, max_len=55)}_{datetime.now():%Y%m%d_%H%M%S}.txt"
    report_path.write_text(report, encoding="utf-8")
    return report, report_path


def scan_report(settings: Settings, project_name: str) -> str:
    meta = load_project_meta(settings, project_name)
    documents, warnings = read_materials(settings, project_name)
    lines = [project_summary(settings, project_name), "", "Читаемые текстовые материалы:"]
    if documents:
        for doc in documents:
            lines.append(f"• {doc.path.name}: {len(doc.text)} символов")
    else:
        lines.append("• нет")
    if warnings:
        lines += ["", "Предупреждения:"] + [f"• {w}" for w in warnings]
    lines += ["", f"project.json schema: {meta.get('schema_version')}"]
    return "\n".join(lines)


def _cli() -> int:
    parser = argparse.ArgumentParser(description="AI Agent V3")
    sub = parser.add_subparsers(dest="command", required=True)
    sub.add_parser("list")
    p_create = sub.add_parser("create")
    p_create.add_argument("name")
    p_create.add_argument("--type", default="other")
    p_scan = sub.add_parser("scan")
    p_scan.add_argument("project")
    p_info = sub.add_parser("info")
    p_info.add_argument("project")
    p_solve = sub.add_parser("solve")
    p_solve.add_argument("project")
    p_solve.add_argument("request")
    p_check = sub.add_parser("check")
    p_check.add_argument("project")
    p_check.add_argument("document")
    p_plan = sub.add_parser("plan")
    p_plan.add_argument("project")
    p_plan.add_argument("--force", action="store_true")
    p_sec = sub.add_parser("section")
    p_sec.add_argument("project")
    p_sec.add_argument("section_id")
    p_sec.add_argument("--instruction", default="")
    p_ch = sub.add_parser("chapter")
    p_ch.add_argument("project")
    p_ch.add_argument("chapter_id")
    p_asm = sub.add_parser("assemble")
    p_asm.add_argument("project")
    p_asm.add_argument("--force", action="store_true")
    args = parser.parse_args()

    try:
        settings = load_settings()
        if args.command == "list":
            projects = list_projects(settings)
            print("\n".join(p.name for p in projects) if projects else "Папок пока нет.")
        elif args.command == "create":
            print(create_project(settings, args.name, args.type))
        elif args.command == "scan":
            print(scan_report(settings, args.project))
        elif args.command == "info":
            print(project_summary(settings, args.project))
            print("\n" + readiness_text(settings, args.project))
            print("\n" + materials_text(settings, args.project))
            plan = load_work_plan(get_project(settings, args.project))
            if plan:
                print("\n" + plan_text(plan))
        elif args.command == "solve":
            result = solve_task(settings, args.project, args.request)
            print(result.answer)
            print(f"DOCX: {result.docx_path}")
        elif args.command == "check":
            report, path = check_formatting(settings, args.project, args.document)
            print(report)
            print(f"\nОтчёт: {path}")
        elif args.command == "plan":
            plan = create_work_plan(settings, args.project, force=args.force)
            print(plan_text(plan))
        elif args.command == "section":
            result = generate_section(settings, args.project, args.section_id, user_instruction=args.instruction)
            print(f"{result.node.id}: v{result.version}, {result.status}, score={result.score}")
            print(result.docx_path)
        elif args.command == "chapter":
            result = generate_chapter(settings, args.project, args.chapter_id)
            print(result.docx_path)
        elif args.command == "assemble":
            result = assemble_work(settings, args.project, allow_unapproved=args.force)
            print(result.docx_path)
            print(result.report_path)
        return 0
    except (AgentError, ProjectError, ProviderError, Exception) as exc:
        # CLI intentionally converts all errors to a non-zero code and readable message.
        logger.exception("CLI error")
        print(f"ОШИБКА: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(_cli())
