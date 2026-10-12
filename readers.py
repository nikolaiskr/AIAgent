from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from docx import Document
from PyPDF2 import PdfReader


TEXT_EXTENSIONS = {".pdf", ".docx", ".txt", ".xlsx", ".md"}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}
UPLOAD_EXTENSIONS = TEXT_EXTENSIONS | IMAGE_EXTENSIONS


class ReaderError(RuntimeError):
    pass


@dataclass(slots=True)
class SourceDocument:
    path: Path
    text: str
    warnings: list[str] = field(default_factory=list)


def _read_txt(path: Path) -> SourceDocument:
    errors: list[str] = []
    for encoding in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            return SourceDocument(path=path, text=path.read_text(encoding=encoding))
        except UnicodeDecodeError as exc:
            errors.append(str(exc))
        except OSError as exc:
            raise ReaderError(f"Не удалось прочитать TXT/MD {path.name}: {exc}") from exc
    return SourceDocument(
        path=path,
        text=path.read_text(encoding="utf-8", errors="replace"),
        warnings=errors,
    )


def _read_pdf(path: Path) -> SourceDocument:
    parts: list[str] = []
    warnings: list[str] = []
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                warnings.append("PDF защищён паролем; часть текста может быть недоступна.")
        for i, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception as exc:
                warnings.append(f"Не удалось извлечь текст со страницы {i}: {exc}")
                text = ""
            parts.append(f"\n--- Страница {i} ---\n{text}")
    except Exception as exc:
        raise ReaderError(f"Не удалось прочитать PDF {path.name}: {exc}") from exc

    full_text = "".join(parts).strip()
    if len("".join(full_text.split())) < 30:
        warnings.append("В PDF почти нет извлекаемого текста. Вероятно, это скан; V3 pipeline при настроенном OpenAI Vision попробует распознать страницы как изображения.")
    return SourceDocument(path=path, text=full_text, warnings=warnings)


def _read_docx(path: Path) -> SourceDocument:
    try:
        doc = Document(str(path))
    except Exception as exc:
        raise ReaderError(f"Не удалось прочитать DOCX {path.name}: {exc}") from exc

    parts: list[str] = []
    for idx, paragraph in enumerate(doc.paragraphs, start=1):
        text = paragraph.text.strip()
        if text:
            parts.append(f"[Абзац {idx}] {text}")
    for table_index, table in enumerate(doc.tables, start=1):
        parts.append(f"\n[Таблица {table_index}]")
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return SourceDocument(path=path, text="\n".join(parts))


def _read_xlsx(path: Path) -> SourceDocument:
    try:
        sheets = pd.read_excel(path, sheet_name=None, dtype=object)
    except Exception as exc:
        raise ReaderError(f"Не удалось прочитать XLSX {path.name}: {exc}") from exc

    parts: list[str] = []
    for sheet_name, frame in sheets.items():
        parts.append(f"\n=== Лист: {sheet_name} ===")
        if frame.empty:
            parts.append("[пустой лист]")
            continue
        clean = frame.astype(object).where(pd.notna(frame), "")
        parts.append(clean.to_csv(index=False, sep="\t"))
    return SourceDocument(path=path, text="\n".join(parts))


def read_source(path: Path) -> SourceDocument:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return _read_txt(path)
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".docx":
        return _read_docx(path)
    if suffix == ".xlsx":
        return _read_xlsx(path)
    if suffix in IMAGE_EXTENSIONS:
        return SourceDocument(
            path=path,
            text="",
            warnings=["Изображение сохранено, но локальное распознавание пока не выполняется. Для анализа понадобится GPT vision."],
        )
    raise ReaderError(f"Формат {suffix or '[без расширения]'} пока не поддерживается: {path.name}")
